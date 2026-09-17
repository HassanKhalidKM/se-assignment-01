import os
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = r"C:\Users\HP\Desktop\SE"
report_dir = os.path.join(base_dir, "report")
screenshots_dir = os.path.join(base_dir, "screenshots")
os.makedirs(report_dir, exist_ok=True)

# -------------------------------------------------------------
# 1. Generate Mermaid-style DevOps Flow Diagram using Matplotlib
# -------------------------------------------------------------
def generate_mermaid_diagram():
    diagram_path = os.path.join(report_dir, "devops_flow_diagram.png")
    
    fig, ax = plt.subplots(figsize=(10, 8), dpi=300)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')
    
    # Background
    fig.patch.set_facecolor('#ffffff')
    ax.set_facecolor('#ffffff')
    
    # Title
    ax.text(50, 96, "DevOps Flow — Continuous Integration & Continuous Deployment", 
            ha='center', va='center', fontsize=14, fontweight='bold', color='#1e293b')
    
    def draw_box(x, y, w, h, title, lines, header_color, bg_color, border_color):
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.5,rounding_size=2",
                                      facecolor=bg_color, edgecolor=border_color, linewidth=2)
        ax.add_patch(rect)
        # Header strip
        header_rect = patches.FancyBboxPatch((x, y + h - 5), w, 5, boxstyle="round,pad=0.2,rounding_size=1",
                                             facecolor=header_color, edgecolor=header_color)
        ax.add_patch(header_rect)
        ax.text(x + w/2, y + h - 2.5, title, ha='center', va='center', 
                fontsize=11, fontweight='bold', color='#ffffff')
        
        # Content
        cur_y = y + h - 8
        for line in lines:
            ax.text(x + 2, cur_y, line, ha='left', va='center', 
                    fontsize=9, color='#334155')
            cur_y -= 3.2

    # Box 1: Local Development
    draw_box(10, 68, 80, 22, 
             "1. LOCAL DEVELOPMENT & VERSION CONTROL",
             ["• Technologies: HTML5 (Semantic), CSS3 (Flexbox/Dark theme), Vanilla JavaScript (ES8)",
              "• Testing: Local execution in browser (index.html, responsive check, localStorage)",
              "• Multi-Stage Commits: Structure -> Core Logic -> UI & CI Configuration"],
             "#4f46e5", "#f5f3ff", "#c7d2fe")
    
    # Arrow 1 -> 2
    ax.annotate('', xy=(50, 56), xytext=(50, 68),
                arrowprops=dict(facecolor='#4f46e5', edgecolor='#4f46e5', width=2, headwidth=8))
    ax.text(52, 62, "git push origin main", fontsize=9, fontweight='semibold', color='#4f46e5')

    # Box 2: GitHub Repository
    draw_box(20, 48, 60, 8,
             "2. GITHUB REPOSITORY (Remote Main Branch)",
             ["• Central Repository: HassanKhalidKM/se-assignment-01"],
             "#0284c7", "#f0f9ff", "#bae6fd")

    # Arrow 2 -> 3
    ax.annotate('', xy=(50, 36), xytext=(50, 48),
                arrowprops=dict(facecolor='#0284c7', edgecolor='#0284c7', width=2, headwidth=8))
    ax.text(52, 42, "triggers GitHub Actions", fontsize=9, fontweight='semibold', color='#0284c7')

    # Box 3: GitHub Actions CI
    draw_box(10, 15, 42, 21,
             "3. CI PIPELINE (.github/workflows/ci.yml)",
             ["• Job: Validate & Lint (ubuntu-latest)",
              "  ✓ Check required files exist",
              "  ✓ HTML5 validation (html5validator)",
              "  ✓ JS linting (jshint via .jshintrc)",
              "• Error detection & fast feedback loop"],
             "#16a34a", "#f0fdf4", "#bbf7d0")

    # Box 4: GitHub Pages CD
    draw_box(56, 15, 34, 21,
             "4. CD DEPLOYMENT (.github/workflows/deploy.yml)",
             ["• Target: GitHub Pages",
              "• Trigger: On successful push",
              "• Actions:",
              "  ✓ Upload pages artifact",
              "  ✓ Deploy to Pages live URL"],
             "#ea580c", "#fff7ed", "#fed7aa")

    # Horizontal Arrow CI -> CD
    ax.annotate('', xy=(56, 25.5), xytext=(52, 25.5),
                arrowprops=dict(facecolor='#16a34a', edgecolor='#16a34a', width=2, headwidth=8))

    # Box 5: Live Production
    rect_prod = patches.FancyBboxPatch((30, 2), 40, 7, boxstyle="round,pad=0.3,rounding_size=1.5",
                                      facecolor="#ecfdf5", edgecolor="#10b981", linewidth=2)
    ax.add_patch(rect_prod)
    ax.text(50, 5.5, "★ LIVE DEPLOYMENT: hassankhalidkm.github.io/se-assignment-01/", 
            ha='center', va='center', fontsize=9.5, fontweight='bold', color='#065f46')

    # Arrows 3 & 4 to Prod
    ax.annotate('', xy=(42, 9), xytext=(31, 15),
                arrowprops=dict(facecolor='#10b981', edgecolor='#10b981', width=1.5, headwidth=6))
    ax.annotate('', xy=(58, 9), xytext=(73, 15),
                arrowprops=dict(facecolor='#10b981', edgecolor='#10b981', width=1.5, headwidth=6))

    plt.tight_layout()
    plt.savefig(diagram_path, bbox_inches='tight', dpi=300)
    plt.close()
    return diagram_path

# -------------------------------------------------------------
# 2. Build Formatted Word Document (.docx)
# -------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_word_report(diagram_path):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Assignment 01: Build and Deploy a Small Application")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("Software Engineering — DevOps & Continuous Integration (CI/CD)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)
    run_sub.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Student Name", "Muhammad Hassan Khalid"),
        ("Roll Number", "FA25-BCS-132"),
        ("Course & Assignment", "Software Engineering — Assignment 01"),
        ("Chosen Application", "To-Do List Application (TaskFlow)"),
        ("GitHub Repository", "https://github.com/HassanKhalidKM/se-assignment-01"),
        ("Live Deployed URL", "https://hassankhalidkm.github.io/se-assignment-01/")
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        set_cell_background(c0, "F1F5F9")
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        set_cell_background(c1, "FFFFFF")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(15, 23, 42)
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_bullet(p, bold_prefix, text):
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run("• " + bold_prefix + ": ")
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)

    # 1. Application Name and Purpose
    add_heading("1. Application Name and Purpose")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("• Application Name: TaskFlow\n"
                  "• Category: To-Do List Application (Option 1 from assignment prompt)\n"
                  "• Purpose: TaskFlow is an intuitive, single-page productivity web application designed to help individuals organize, track, and complete daily tasks. Developed with vanilla web standards (HTML5, CSS3, ES8 JavaScript), it requires no external libraries or build-step dependencies, making it ultra-lightweight, resilient, and instantly accessible on desktop and mobile browsers.")
    r.font.size = Pt(10.5)

    # 2. Main Features
    add_heading("2. Main Features")
    features = [
        ("Task Creation", "Quickly create tasks with real-time feedback via text input and Enter key or Add button."),
        ("Status Toggling", "Mark tasks as completed or active with instant animated strike-through styling and checkbox indicators."),
        ("Task Filtering", "Dynamic filter tabs allow users to view 'All', 'Active' (incomplete), or 'Completed' tasks."),
        ("Persistent Storage", "All tasks and their status automatically sync with the browser's localStorage API, ensuring data survives page reloads."),
        ("Live Metrics Badge", "The application header displays a dynamic counter showing remaining pending tasks."),
        ("Accessibility & Fluid Design", "Incorporates semantic HTML5 landmarks, ARIA roles/labels for assistive technologies, and a responsive dark theme.")
    ]
    for bold_pre, desc in features:
        p = doc.add_paragraph()
        add_bullet(p, bold_pre, desc)

    # 3. DevOps Flow Followed
    add_heading("3. DevOps Flow Followed")
    p = doc.add_paragraph()
    r = p.add_run("The project strictly adhered to modern DevOps principles by automating the build, validation, and deployment pipeline using Git and GitHub Actions. Every change pushed to the main branch triggers automated quality gates before proceeding to production deployment.")
    r.font.size = Pt(10.5)

    # Embed Mermaid Diagram
    if os.path.exists(diagram_path):
        p_diag = doc.add_paragraph()
        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_diag.paragraph_format.space_before = Pt(8)
        p_diag.paragraph_format.space_after = Pt(4)
        run_img = p_diag.add_run()
        run_img.add_picture(diagram_path, width=Inches(6.2))
        
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run("Figure 1: Automated DevOps Pipeline Architecture (Continuous Integration & Continuous Deployment)")
        cap_run.font.size = Pt(9)
        cap_run.font.italic = True
        cap_run.font.color.rgb = RGBColor(100, 116, 139)

    # 4. Problems Faced and How They Were Solved
    add_heading("4. Problems Faced and How They Were Solved")
    problems = [
        ("Linter False Positives (JSHint)", 
         "Running jshint with overly strict default command-line flags flagged standard browser runtime variables (such as localStorage and document) as undefined. This was resolved by creating a dedicated .jshintrc configuration file specifying esversion: 8, browser: true, and registering required globals."),
        ("Deliberate Break & Fix Cycle (Step 9)", 
         "To rigorously demonstrate CI failure without permanently breaking the repository, an intentional syntax error (var broken = (;) was introduced and pushed to trigger a failed CI run. Once evidence was captured, the syntax was restored and pushed to produce a green verified build."),
        ("GitHub Pages Deployment Workflow Setup", 
         "By default, GitHub Pages expects deployment from a dedicated branch. To deploy directly from GitHub Actions artifacts, repository settings were updated to set Source to 'GitHub Actions', and permissions (pages: write, id-token: write) were granted in deploy.yml.")
    ]
    for bold_pre, desc in problems:
        p = doc.add_paragraph()
        add_bullet(p, bold_pre, desc)

    # 5. What You Learned from Continuous Integration (CI)
    add_heading("5. What You Learned from Continuous Integration (CI)")
    learnings = [
        ("Automated Quality Assurance", "CI acts as an impartial safety net. Syntax errors, missing files, or bad markup are trapped immediately before impacting end users."),
        ("Fast Feedback Loop", "Developers receive automated failure reports within seconds of pushing, pinpointing exact files and lines containing regressions."),
        ("Deterministic Environments", "Running tests on standardized Ubuntu virtual environments eliminates discrepancies caused by differing developer operating systems."),
        ("Seamless CI/CD Hand-off", "Successfully passing CI builds can automatically feed directly into CD deployment pipelines (such as GitHub Pages), eliminating manual deployment friction.")
    ]
    for bold_pre, desc in learnings:
        p = doc.add_paragraph()
        add_bullet(p, bold_pre, desc)

    # 6. Deliverable Screenshots
    add_heading("6. Required Deliverables & Screenshots")

    screenshot_items = [
        ("app_local.png", "Screenshot 1: Running Application (Local browser verification)"),
        ("commits.png", "Screenshot 2: Git Commit History (Multi-stage meaningful commits)"),
        ("ci_failed.png", "Screenshot 3: Failed CI Workflow (Intentional syntax error caught by GitHub Actions)"),
        ("ci_success.png", "Screenshot 4: Successful CI Workflow (Clean build after applying fix)"),
        ("pages_live.png", "Screenshot 5: Deployed Application (Live on GitHub Pages)")
    ]

    for fname, caption_text in screenshot_items:
        fpath = os.path.join(screenshots_dir, fname)
        if os.path.exists(fpath):
            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(10)
            p_head.paragraph_format.space_after = Pt(4)
            r_head = p_head.add_run(caption_text)
            r_head.font.bold = True
            r_head.font.size = Pt(11)

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_pic = p_img.add_run()
            r_pic.add_picture(fpath, width=Inches(5.8))
            
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_after = Pt(6)

    docx_path = os.path.join(report_dir, "MuhammadHassanKhalid_FA25-BCS-132_SEAss01_Report.docx")
    doc.save(docx_path)
    return docx_path

if __name__ == "__main__":
    diag = generate_mermaid_diagram()
    docx = create_word_report(diag)
    print(f"SUCCESS: Diagram saved to {diag}")
    print(f"SUCCESS: Word Document saved to {docx}")
